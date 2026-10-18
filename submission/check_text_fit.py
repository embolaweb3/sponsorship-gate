#!/usr/bin/env python3
"""Stricter text-fit check for a generated deck, using real font metrics.

verify_deck.py's heuristic guesses characters-per-line. This one wraps each
paragraph with the actual Arial / Consolas glyph widths (read from the Windows
fonts folder via WSL) and compares the wrapped height to the shape's height.
It is a check, not a renderer: run it, then eyeball the deck once.

  python3 check_text_fit.py Sponsorship-Gate-Wave2-Pitch-Deck.pptx
"""
import sys

from PIL import ImageFont
from pptx import Presentation

FONT_DIRS = ["/mnt/c/Windows/Fonts/", "/usr/share/fonts/truetype/dejavu/"]
FILES = {
    ("Arial", False): ["arial.ttf", "DejaVuSans.ttf"],
    ("Arial", True): ["arialbd.ttf", "DejaVuSans-Bold.ttf"],
    ("Consolas", False): ["consola.ttf", "DejaVuSansMono.ttf"],
    ("Consolas", True): ["consolab.ttf", "DejaVuSansMono-Bold.ttf"],
}
_cache = {}


def font(name, bold, size_pt):
    key = (name if name in ("Arial", "Consolas") else "Arial", bool(bold))
    k = (key, round(size_pt * 4))
    if k in _cache:
        return _cache[k]
    for d in FONT_DIRS:
        for f in FILES[key]:
            try:
                ft = ImageFont.truetype(d + f, int(round(size_pt * 4)))  # 4 px per pt for precision
                _cache[k] = ft
                return ft
            except OSError:
                continue
    raise RuntimeError("no font found for " + str(key))


def wrapped_lines(text, ft, width_px):
    lines = 0
    for para in text.split("\n"):
        words = para.split(" ")
        cur = ""
        n = 1
        for w in words:
            trial = (cur + " " + w).strip()
            if ft.getlength(trial) <= width_px or not cur:
                cur = trial
            else:
                n += 1
                cur = w
        lines += n
    return lines


def main(path):
    prs = Presentation(path)
    bad = 0
    for si, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not shp.has_text_frame or not shp.text_frame.text.strip():
                continue
            tf = shp.text_frame
            ml = (tf.margin_left or 0) + (tf.margin_right or 0)
            width_pt = (shp.width - ml) / 12700
            need_pt = 0.0
            for p in tf.paragraphs:
                if not p.runs:
                    continue
                r = p.runs[0]
                size = r.font.size.pt if r.font.size else 18
                ft = font(r.font.name or "Arial", r.font.bold, size)
                text = "".join(x.text for x in p.runs)
                n = wrapped_lines(text, ft, width_pt * 4)
                spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                after = p.space_after.pt if p.space_after is not None else 0
                need_pt += n * size * 1.2 * spacing + after
            have_pt = (shp.height - (tf.margin_top or 0) - (tf.margin_bottom or 0)) / 12700
            if need_pt > have_pt + 3:
                bad += 1
                print(f"slide {si}: '{tf.text[:46]!r}' needs {need_pt:.0f}pt, box {have_pt:.0f}pt")
    print("text-fit:", "OK" if not bad else f"{bad} box(es) overflow")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
